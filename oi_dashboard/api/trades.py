"""Paper-trade API: listing, manual/bulk entry, reset, exports and tearsheet."""

from datetime import datetime

from flask import Blueprint, jsonify, request, send_file

from oi_dashboard.core import state
from oi_dashboard.core.constants import DEFAULT_LOTS, IST, LOT_SIZE
from oi_dashboard.storage.session_store import save_paper_trades
from oi_dashboard.trading.accounting import (
    _dte_adjusted_stop_mult,
    check_gap_reversal_exit,
    compute_trade_costs,
    target_decay_for_dte,
)
from oi_dashboard.trading.fills import get_leg_price

bp = Blueprint("trades", __name__)


@bp.route("/api/paper_trades")
def api_paper_trades():
    return jsonify({
        "trades": state.paper_trades,
        "date":   state.paper_trade_date,
        "count":  len(state.paper_trades),
        "open":   sum(1 for t in state.paper_trades if t["status"] == "OPEN"),
    })


@bp.route("/api/auto/kill", methods=["POST"])
def api_auto_kill():
    try:
        now_str = datetime.now(IST).strftime("%H:%M:%S")
        killed  = 0
        for trade in state.paper_trades:
            if trade.get("status") != "OPEN":
                continue
            if trade.get("trade_type") == "IRON_CONDOR":
                net_cur  = trade.get("net_credit_current", trade.get("net_credit_entry", 0))
                pnl_unit = trade.get("net_credit_entry", 0) - net_cur
                lots     = trade.get("lots", DEFAULT_LOTS)
                trade.update(status="KILLED", exit_time=now_str,
                             exit_reason="MANUAL_KILL",
                             pnl_per_lot=round(pnl_unit * LOT_SIZE, 2),
                             pnl_total=round(pnl_unit * LOT_SIZE * lots, 2))
            elif trade.get("trade_type") in ("STRANGLE", "GAP_SPREAD"):
                net_cur  = trade.get("net_credit_current", trade.get("net_credit_entry", 0))
                pnl_unit = trade.get("net_credit_entry", 0) - net_cur
                lots     = trade.get("lots", DEFAULT_LOTS)
                trade.update(status="KILLED", exit_time=now_str,
                             exit_reason="MANUAL_KILL",
                             pnl_per_lot=round(pnl_unit * LOT_SIZE, 2),
                             pnl_total=round(pnl_unit * LOT_SIZE * lots, 2))
            elif trade.get("trade_type") == "PANIC_BUY":
                cur      = trade.get("current_ltp", trade.get("entry_ltp", 0))
                lots     = trade.get("lots", 3)
                pnl_unit = cur - trade.get("entry_ltp", 0)
                trade.update(status="KILLED", exit_time=now_str,
                             exit_reason="MANUAL_KILL",
                             exit_ltp=round(cur, 2),
                             pnl_per_lot=round(pnl_unit * LOT_SIZE, 2),
                             pnl_total=round(pnl_unit * LOT_SIZE * lots, 2))
                state.panic_buy_active = False
            else:
                cur = trade.get("current_ltp", trade.get("entry_ltp", 0))
                trade.update(status="KILLED", exit_time=now_str,
                             exit_reason="MANUAL_KILL",
                             pnl_lot=round((trade.get("entry_ltp", 0) - cur) * LOT_SIZE, 2))
            killed += 1
        save_paper_trades()
        return jsonify({"ok": True, "killed": killed})
    except Exception as e:
        print(f"[Kill] Error: {e}")
        return jsonify({"ok": False, "error": str(e)}), 200


@bp.route("/api/paper_trades/reset")
def api_paper_trades_reset():
    try:
        state.paper_trades          = []
        state.paper_trade_ids       = set()
        state.paper_trade_paused    = False
        state.ic_entry_date         = None
        state.panic_buy_active      = False
        state.panic_buy_entry_time  = None
        state.panic_buy_date        = None
        state.spot_10min_history    = []
        state.spot_60min_history    = []
        state.strike_session_lows   = {}
        state.cache_time            = 0
        save_paper_trades()
        return jsonify({"reset": True})
    except Exception as e:
        print(f"[Reset] Error: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"reset": False, "error": str(e)})


@bp.route("/api/paper_trades/force_reversal_check", methods=["POST"])
def api_force_reversal_check():
    """Manually trigger gap reversal exit on all open GAP_SPREAD trades."""
    try:
        spot = state.last_known_spot or 0
        if spot < 1000:
            return jsonify({"ok": False, "error": "no valid spot price available"})
        time_str = datetime.now(IST).strftime("%H:%M:%S")
        before_open = [t for t in state.paper_trades if t.get("status") == "OPEN"
                       and t.get("trade_type") == "GAP_SPREAD"]
        check_gap_reversal_exit(spot, time_str)
        closed = [t for t in before_open if t.get("status") != "OPEN"]
        return jsonify({
            "ok":       True,
            "spot":     spot,
            "checked":  len(before_open),
            "closed":   len(closed),
            "exits":    [{"id": t["id"], "reason": t.get("exit_reason")} for t in closed],
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(e)})


@bp.route("/api/paper_trades/resume")
def api_paper_trades_resume():
    state.paper_trade_paused = False
    return jsonify({"paused": False})


@bp.route("/api/paper_trades/manual_entry", methods=["POST"])
def api_manual_entry():
    """Manual paper trade entry — bypasses can_enter() checks, marks manual_entry=True."""
    try:
        d        = request.get_json() or {}
        tt       = d.get("trade_type", "IRON_CONDOR")
        lots     = int(d.get("lots", DEFAULT_LOTS))
        spot     = state.last_known_spot or 0
        time_str = datetime.now(IST).strftime("%H:%M:%S")

        chain    = (state.cache or {}).get("chain", [])
        cmap     = {r["strike"]: r for r in chain}

        def _leg_price(strike, side, action):
            row = cmap.get(int(strike)) or {}
            p   = get_leg_price(row, side.lower(), action)
            if p is not None:
                return p
            # Manual override from form if API has no data for this strike
            ltp_override = float(d.get(f"{side.lower()}_ltp_{strike}", 0) or 0)
            if ltp_override > 0:
                return ltp_override
            return None

        if tt == "IRON_CONDOR":
            pe_sell = int(d["pe_sell"]); pe_buy = int(d["pe_buy"])
            ce_sell = int(d["ce_sell"]); ce_buy = int(d["ce_buy"])
            pe_sl = _leg_price(pe_sell, "PE", "sell"); pe_bl = _leg_price(pe_buy, "PE", "buy")
            ce_sl = _leg_price(ce_sell, "CE", "sell"); ce_bl = _leg_price(ce_buy, "CE", "buy")
            if None in (pe_sl, pe_bl, ce_sl, ce_bl):
                return jsonify({"ok": False, "error": f"Untradeable leg: PE_sell={pe_sl} PE_buy={pe_bl} CE_sell={ce_sl} CE_buy={ce_bl}"})
            net   = round(pe_sl - pe_bl + ce_sl - ce_bl, 2)
            trade = {
                "id": len(state.paper_trades) + 1, "trade_type": "IRON_CONDOR",
                "entry_time": time_str, "entry_spot": round(spot, 2), "lots": lots,
                "pe_sell_strike": pe_sell, "pe_sell_entry": pe_sl, "pe_sell_current": pe_sl,
                "pe_buy_strike":  pe_buy,  "pe_buy_entry":  pe_bl, "pe_buy_current":  pe_bl,
                "ce_sell_strike": ce_sell, "ce_sell_entry": ce_sl, "ce_sell_current": ce_sl,
                "ce_buy_strike":  ce_buy,  "ce_buy_entry":  ce_bl, "ce_buy_current":  ce_bl,
                "net_credit_entry": net, "net_credit_current": net,
                "target_net": round(net * 0.40, 2), "stop_net": round(net * 2.0, 2),
                "status": "OPEN", "exit_time": None, "exit_reason": None,
                "pnl_per_lot": 0.0, "pnl_total": 0.0, "carry_forward": False,
                "pe_status": "OPEN", "ce_status": "OPEN",
                "panic_triggered": False, "manual_entry": True, "entry_reason": "manual", "lot_size_used": LOT_SIZE,
                "insurance_cost_lot": round(pe_bl + ce_bl, 2),
                "insurance_cost_total": round((pe_bl + ce_bl) * LOT_SIZE * lots, 2),
            }
        elif tt == "GAP_SPREAD":
            side    = d.get("side", "CE")
            sell_s  = int(d["sell_strike"]); buy_s = int(d["buy_strike"])
            sl      = _leg_price(sell_s, side, "sell"); bl = _leg_price(buy_s, side, "buy")
            if None in (sl, bl):
                return jsonify({"ok": False, "error": f"Untradeable leg: sell {sell_s}={sl} buy {buy_s}={bl}"})
            net     = round(sl - bl, 2)
            direction = d.get("direction", "GAP_DOWN")
            trade = {
                "id": len(state.paper_trades) + 1, "trade_type": "GAP_SPREAD",
                "direction": direction, "side": side,
                "entry_time": time_str, "entry_spot": round(spot, 2), "lots": lots,
                "sell_strike": sell_s, "sell_entry": sl, "sell_current": sl,
                "buy_strike":  buy_s,  "buy_entry":  bl, "buy_current":  bl,
                "net_credit_entry": net, "net_credit_current": net,
                "target_net": round(net * 0.40, 2), "stop_net": round(net * 2.0, 2),
                "status": "OPEN", "exit_time": None, "exit_reason": None,
                "pnl_per_lot": 0.0, "pnl_total": 0.0, "max_adverse_spot": 999.0,
                "carry_forward": False, "prev_close": 0.0, "open_price": 0.0,
                "manual_entry": True, "entry_reason": "manual", "lot_size_used": LOT_SIZE,
            }
        else:
            return jsonify({"ok": False, "error": f"unsupported trade_type: {tt}"})

        state.paper_trades.append(trade)
        save_paper_trades()
        print(f"[MANUAL] {tt} #{trade['id']} lots={lots} net=₹{trade.get('net_credit_entry',0):.2f}")
        return jsonify({"ok": True, "trade_id": trade["id"]})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(e)})


@bp.route("/api/paper_trades/add", methods=["POST"])
def api_paper_trades_add():
    try:
        data       = request.get_json()
        strike     = int(data["strike"])
        otype      = str(data["option_type"]).upper()
        entry_time = str(data["entry_time"])
        entry_ltp  = float(data["entry_ltp"])
        entry_spot = float(data.get("entry_spot", 0))
        entry_ovr  = float(data.get("entry_ovr", 0))
        entry_fair = float(data.get("entry_fair", 0))

        if (strike, otype) in state.paper_trade_ids:
            return jsonify({"error": f"{strike} {otype} already exists in today's trades"}), 400

        if not state.paper_trade_date:
            state.paper_trade_date = datetime.now(IST).strftime("%Y-%m-%d")

        state.paper_trade_ids.add((strike, otype))
        state.paper_trades.append({
            "id":               len(state.paper_trades) + 1,
            "strike":           strike,
            "option_type":      otype,
            "entry_time":       entry_time,
            "entry_ltp":        entry_ltp,
            "entry_fair":       entry_fair,
            "entry_ovr":        entry_ovr,
            "entry_spot":       entry_spot,
            "stop_price":       round(entry_ltp * _dte_adjusted_stop_mult(5), 2),   # 5 = mid-week default for manual
            "target_price":     round(entry_ltp * (1 - target_decay_for_dte(5)), 2),
            "current_ltp":      entry_ltp,
            "current_ovr":      entry_ovr,
            "current_spot":     entry_spot,
            "max_ltp":          entry_ltp,
            "min_ltp":          entry_ltp,
            "max_adverse_spot": 0.0,
            "status":           "OPEN",
            "exit_time":        None,
            "exit_ltp":         None,
            "exit_reason":      None,
            "pnl_unit":         0.0,
            "pnl_lot":          0.0,
            "updates":          [],
            "manual_entry":     True,
            "entry_reason":     "manual",
            "lot_size_used":     LOT_SIZE,
        })
        save_paper_trades()
        return jsonify({
            "success":  True,
            "trade_id": len(state.paper_trades),
            "message":  f"Added {strike} {otype} @ ₹{entry_ltp} at {entry_time}",
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 400


@bp.route("/api/paper_trades/bulk_add", methods=["POST"])
def api_paper_trades_bulk_add():
    try:
        items   = request.get_json()
        if not isinstance(items, list):
            return jsonify({"error": "Expected JSON array"}), 400

        added   = 0
        skipped = 0
        today   = datetime.now(IST).strftime("%Y-%m-%d")
        if not state.paper_trade_date:
            state.paper_trade_date = today

        for data in items:
            strike     = int(data["strike"])
            otype      = str(data["option_type"]).upper()
            entry_time = str(data.get("entry_time", "09:15:00"))
            entry_ltp  = float(data["entry_ltp"])
            entry_spot = float(data.get("entry_spot", 0))
            entry_ovr  = float(data.get("entry_ovr", 0))
            entry_fair = float(data.get("entry_fair", 0))

            if (strike, otype) in state.paper_trade_ids:
                skipped += 1
                continue

            state.paper_trade_ids.add((strike, otype))
            state.paper_trades.append({
                "id":               len(state.paper_trades) + 1,
                "strike":           strike,
                "option_type":      otype,
                "entry_time":       entry_time,
                "entry_ltp":        entry_ltp,
                "entry_fair":       entry_fair,
                "entry_ovr":        entry_ovr,
                "entry_spot":       entry_spot,
                "stop_price":       round(entry_ltp * _dte_adjusted_stop_mult(5), 2),
                "target_price":     round(entry_ltp * (1 - target_decay_for_dte(5)), 2),
                "current_ltp":      entry_ltp,
                "current_ovr":      entry_ovr,
                "current_spot":     entry_spot,
                "max_ltp":          entry_ltp,
                "min_ltp":          entry_ltp,
                "max_adverse_spot": 0.0,
                "status":           "OPEN",
                "exit_time":        None,
                "exit_ltp":         None,
                "exit_reason":      None,
                "pnl_unit":         0.0,
                "pnl_lot":          0.0,
                "updates":          [],
                "manual_entry":     True,
                "entry_reason":     "manual",
                "lot_size_used":     LOT_SIZE,
            })
            added += 1

        state.paper_trade_paused = False   # resume auto-detection after bulk load
        state.cache_time = 0               # bust cache so next /api/data returns fresh paper_trades
        save_paper_trades()
        return jsonify({
            "success": True,
            "added":   added,
            "skipped": skipped,
            "total":   len(state.paper_trades),
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 400


@bp.route("/api/paper_trades/excel")
def api_paper_trades_excel():
    import io
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return jsonify({"error": "openpyxl not installed — run: pip install openpyxl"}), 500

    try:
        wb       = openpyxl.Workbook()
        hdr_fill = PatternFill("solid", fgColor="1F2937")
        hdr_font = Font(bold=True, color="FFFFFF")
        ctr      = Alignment(horizontal="center")

        # Sheet 1 — Trade Summary (works for GAP_SPREAD + single-leg)
        ws1 = wb.active
        ws1.title = "Trade Summary"
        cols1 = ["ID","Trade Type","Direction","Side/Type","Sell Strike","Buy Strike",
                 "Net Credit Entry","Net Credit Exit","Entry Time","Entry Spot","Lots","DTE",
                 "Exit Time","Exit Reason","P&L/Lot","P&L Total","Costs","P&L Net","Max Adverse Spot","Status","Date"]
        for c, h in enumerate(cols1, 1):
            cell = ws1.cell(row=1, column=c, value=h)
            cell.fill, cell.font, cell.alignment = hdr_fill, hdr_font, ctr

        status_colors = {"OPEN":"0ea5e9","EOD":"f59e0b","KILLED":"ef4444",
                         "TARGET_HIT":"22c55e","STOPPED":"ef4444","EOD_EXIT":"f59e0b"}

        for ri, t in enumerate(state.paper_trades, 2):
            tt   = t.get("trade_type", "SINGLE")
            pnl  = t.get("pnl_total", (t.get("pnl_lot") or t.get("pnl_per_lot", 0)))
            compute_trade_costs(t)
            costs   = t.get("costs_total", "--")
            pnl_net = t.get("pnl_net", "--")
            vals = [
                t.get("id", ri-1),
                tt,
                t.get("direction", "--"),
                t.get("side", t.get("option_type", "--")),
                t.get("sell_strike", t.get("strike", "--")),
                t.get("buy_strike", "--"),
                t.get("net_credit_entry", t.get("entry_ltp", "--")),
                t.get("net_credit_current", t.get("current_ltp", "--")),
                t.get("entry_time", "--"),
                t.get("entry_spot", "--"),
                t.get("lots", 1),
                t.get("dte_at_entry", "--"),
                t.get("exit_time", "--"),
                t.get("exit_reason", t.get("status", "--")),
                t.get("pnl_per_lot", "--"),
                pnl,
                costs,
                pnl_net,
                t.get("max_adverse_spot", "--"),
                t.get("status", "--"),
                state.paper_trade_date,
            ]
            for c, v in enumerate(vals, 1):
                cell = ws1.cell(row=ri, column=c, value=v)
                if c == 16 and isinstance(v, (int, float)):
                    cell.font = Font(color="22c55e" if v >= 0 else "ef4444", bold=True)
                if c == 17 and isinstance(v, (int, float)):  # Costs — always negative-ish
                    cell.font = Font(color="f59e0b", bold=True)
                if c == 18 and isinstance(v, (int, float)):  # P&L Net
                    cell.font = Font(color="22c55e" if v >= 0 else "ef4444", bold=True)
                if c == 20:
                    cell.font = Font(color=status_colors.get(str(v), "ffffff"))

        for col in ws1.columns:
            ws1.column_dimensions[col[0].column_letter].width = min(
                max(len(str(c.value or "")) for c in col) + 3, 25)

        # Sheet 2 — GAP_SPREAD details
        ws2 = wb.create_sheet("Spread Details")
        ws2.append(["ID","Type","Sell Strike","Sell Entry","Sell Exit","Buy Strike",
                    "Buy Entry","Buy Exit","Net Entry","Net Exit","Credit Kept","P&L/Lot","P&L Total","Lots"])
        for t in state.paper_trades:
            if t.get("trade_type") in ("GAP_SPREAD", "IRON_CONDOR", "STRANGLE"):
                ws2.append([
                    t.get("id"), t.get("trade_type"),
                    t.get("sell_strike", "--"), t.get("sell_entry", "--"), t.get("sell_current", "--"),
                    t.get("buy_strike", "--"),  t.get("buy_entry",  "--"), t.get("buy_current", "--"),
                    t.get("net_credit_entry", "--"), t.get("net_credit_current", "--"),
                    round((t.get("net_credit_entry",0) or 0) - (t.get("net_credit_current",0) or 0), 2),
                    t.get("pnl_per_lot", "--"), t.get("pnl_total", "--"), t.get("lots", "--"),
                ])

        # Sheet 3 — Price History (single-leg only; GAP_SPREAD updates stored differently)
        ws3 = wb.create_sheet("Price History")
        ws3.append(["Trade ID","Strike","Type","Time","LTP","OVR%","Spot"])
        for t in state.paper_trades:
            for u in t.get("updates", []):
                ws3.append([t.get("id"), t.get("strike", t.get("sell_strike", "--")),
                             t.get("option_type", t.get("trade_type", "--")),
                             u.get("time"), u.get("ltp"), u.get("ovr"), u.get("spot")])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        from flask import send_file
        return send_file(buf, as_attachment=True,
                         download_name=f"paper_trades_{state.paper_trade_date}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        import traceback
        print(f"[EXCEL] Error: {e}")
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@bp.route("/api/paper_trades/flag_fake_hedges", methods=["POST"])
def api_flag_fake_hedges():
    """
    Step 7: Iterate all in-memory trades and flag any where hedge legs were recorded
    at exactly ₹0.50 (the old hardcoded fallback). Flagged trades are excluded from
    session P&L and marked for audit.
    """
    flagged = 0
    for t in state.paper_trades:
        tt = t.get("trade_type", "")
        if tt == "IRON_CONDOR":
            if t.get("pe_buy_entry") == 0.50 or t.get("ce_buy_entry") == 0.50:
                t["audit_flag"]            = "FAKE_HEDGE_050"
                t["real_pnl_unreliable"]   = True
                t["pnl_excluded_from_session"] = True
                flagged += 1
        elif tt == "GAP_SPREAD":
            if t.get("buy_entry") == 0.50:
                t["audit_flag"]            = "FAKE_HEDGE_050"
                t["real_pnl_unreliable"]   = True
                t["pnl_excluded_from_session"] = True
                flagged += 1
    if flagged:
        save_paper_trades()
    return jsonify({"ok": True, "flagged": flagged, "total": len(state.paper_trades)})


@bp.route("/api/paper/tearsheet")
def api_paper_tearsheet():
    """GET /api/paper/tearsheet?days=180 — build and download tearsheet Excel."""
    try:
        import os
        import tempfile

        from oi_dashboard.trading import paper_engine as PE
        days = max(7, min(int(request.args.get("days", 180)), 730))
        tmp  = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.close()
        PE.export_tearsheet(tmp.name, days=days)
        from datetime import datetime as _dt
        fname = f"tearsheet_{_dt.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return send_file(tmp.name, as_attachment=True,
                         download_name=fname,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except ImportError as e:
        return jsonify({"status": "error", "error": str(e),
                        "hint": "pip install openpyxl"}), 500
    except Exception as e:
        import traceback as _tb
        return jsonify({"status": "error", "error": str(e),
                        "traceback": _tb.format_exc()}), 500
